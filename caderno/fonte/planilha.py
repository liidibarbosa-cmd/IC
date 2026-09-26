"""Planilha de quantitativos (XLSX) com fórmulas: pisos, memória de rodapés, vãos e transições."""
import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

AQUI = os.path.dirname(os.path.abspath(__file__))
D = json.load(open(os.path.join(AQUI, 'dados.json')))
OUT = os.path.join(AQUI, '..', 'quantitativos', 'QUANTITATIVOS PISOS E RODAPES - R00.xlsx')
HEAD = PatternFill('solid', fgColor='3F422F'); SEC = PatternFill('solid', fgColor='ECE4D4'); SUB = PatternFill('solid', fgColor='F1EAD9')
HF = Font(bold=True, color='F5E7C4'); B = Font(bold=True); SF = Font(bold=True, color='7A3423')
thin = Border(bottom=Side(style='thin', color='D8CFBC'))
wb = Workbook()

def header(ws, cols, widths):
    ws.append(cols)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        c = ws.cell(row=ws.max_row, column=i); c.fill = HEAD; c.font = HF; c.alignment = Alignment(vertical='center', wrap_text=True)
    ws.freeze_panes = 'A2'

# --- Pisos
ws = wb.active; ws.title = 'Pisos'
header(ws, ['Código', 'Ambiente', 'Produto', 'Fabricante', 'Formato', 'Acabamento', 'Área adotada (m²)', 'Área geometria DWG (m²)', 'Fonte da medida', 'Controle'],
       [9, 42, 38, 14, 14, 20, 16, 18, 90, 40])
subs = {}
for cod in ('P01', 'P02', 'PP01', 'NE'):
    ws.append([cod, {'NE': 'PISO NÃO DEFINIDO — A CONFIRMAR'}.get(cod, D['produtos'][cod]['produto'].upper())])
    for c in ws[ws.max_row]: c.fill = SEC; c.font = SF
    r0 = ws.max_row + 1
    for p in [x for x in D['pisos'] if x['cod'] == cod]:
        pr = D['produtos'][cod]
        ws.append([cod, p['ambiente'], pr['produto'], pr['fabricante'], pr['formato'], pr['acabamento'], p['area'],
                   round(sum(D['zonas'][z]['area_geo'] for z in p['zonas']), 3), p['fonte'], ' + '.join(p['controle'])])
    r1 = ws.max_row
    ws.append(['', f'Subtotal {cod}', '', '', '', '', f'=SUM(G{r0}:G{r1})'])
    subs[cod] = ws.max_row
    for c in ws[ws.max_row]: c.fill = SUB; c.font = B
ws.append(['', 'TOTAL GERAL ESPECIFICADO (P01 + P02 + PP01)', '', '', '', '', f"=G{subs['P01']}+G{subs['P02']}+G{subs['PP01']}"])
for c in ws[ws.max_row]: c.font = B
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.border = thin; c.alignment = Alignment(vertical='top', wrap_text=True)
    row[6].number_format = '0.00'; row[7].number_format = '0.000'

# --- Rodapés (memória)
ws = wb.create_sheet('Rodapés - memória')
header(ws, ['Código', 'Ambiente', 'Tipo', 'Descrição', 'X início', 'Y início', 'X fim', 'Y fim', 'Comprimento (m)', 'Fonte / observação'],
       [9, 34, 18, 44, 10, 10, 10, 10, 16, 60])
tot = {}
for cod in ('R01', 'R02'):
    ws.append([cod, {'R01': 'RODAPÉ SANTORINI OFF WHITE PADRÃO — h = 8 cm', 'R02': 'RODAPÉ SANTORINI OFF WHITE HARD — h = 8 cm'}[cod]])
    for c in ws[ws.max_row]: c.fill = SEC; c.font = SF
    liq_rows = []
    for r in [x for x in D['rodapes'] if x['cod'] == cod]:
        a = ws.max_row + 1
        for t in r['trechos']:
            ws.append([cod, r['ambiente'], 'trecho de parede', t['nome'], t['a'][0], t['a'][1], t['b'][0], t['b'][1], t['L'], 'geometria DWG (face interna)'])
        b = ws.max_row
        ws.append([cod, r['ambiente'], 'soma dos trechos', '', '', '', '', '', f'=SUM(I{a}:I{b})'])
        s = ws.max_row
        va = ws.max_row + 1
        for v in r['vaos']:
            V = D['vaos'][v['id']]
            ws.append([cod, r['ambiente'], 'vão descontado', f"{v['cod']} — {V['amb']} (trecho {v['trecho']})", V['a'][0], V['a'][1], V['b'][0], V['b'][1], v['larg'], V['origem']])
        vb = ws.max_row
        ws.append([cod, r['ambiente'], 'LÍQUIDO', 'soma dos trechos − vãos', '', '', '', '', f'=I{s}-SUM(I{va}:I{vb})', r['obs']])
        liq_rows.append(ws.max_row)
        for c in ws[ws.max_row]: c.font = B
    ws.append(['', f'Subtotal {cod}', '', '', '', '', '', '', '=' + '+'.join(f'I{x}' for x in liq_rows)])
    tot[cod] = ws.max_row
    for c in ws[ws.max_row]: c.fill = SUB; c.font = B
ws.append(['', 'TOTAL GERAL DE RODAPÉS (R01 + R02)', '', '', '', '', '', '', f"=I{tot['R01']}+I{tot['R02']}"])
for c in ws[ws.max_row]: c.font = B
for row in ws.iter_rows(min_row=2):
    for c in row: c.border = thin; c.alignment = Alignment(vertical='top', wrap_text=True)
    for k in (4, 5, 6, 7, 8): row[k].number_format = '0.000'

# --- Vãos
ws = wb.create_sheet('Vãos')
header(ws, ['ID', 'Código', 'Ambientes', 'Vão DWG (m)', 'Quadro de Esquadrias Rev. 01 (m)', 'Diferença (m)', 'Origem da medida DWG'], [10, 9, 36, 13, 18, 13, 44])
for vid, v in D['vaos'].items():
    ws.append([vid, v['cod'], v['amb'], v['larg'], D['caderno_larg'][v['cod']], f'=E{ws.max_row + 1}-D{ws.max_row + 1}', v['origem']])

# --- Transições
ws = wb.create_sheet('Transições')
header(ws, ['Código', 'Porta', 'Encontro de pisos', 'Situação / decisão', 'Controle'], [8, 12, 50, 80, 14])
for t in D['transicoes']:
    ws.append(list(t))
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(vertical='top', wrap_text=True)
wb.save(OUT)
print('ok', OUT)

"""Lista técnica para orçamento — revestimentos, louças e metais (compra na Bandini).

Quantidades de revestimento vêm dos cadernos 01/02 (dados.json) e 02/02 (dados_paredes.json);
louças e metais vêm da "Lista técnica de compras — Residência IC" (25/09/2026), sem alteração.
Quantidades líquidas: sem perdas de compra ou de corte.
"""
import json, math, os, html
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.drawing.image import Image as XLImage

AQUI = os.path.dirname(os.path.abspath(__file__))
D = json.load(open(os.path.join(AQUI, 'dados.json')))
DP = json.load(open(os.path.join(AQUI, 'dados_paredes.json')))
DATA, REVN = '26/09/2026', '00'
esc = lambda s: html.escape(str(s))
br = lambda v, n=2: f'{v:,.{n}f}'.replace(',', 'X').replace('.', ',').replace('X', '.')

# ------------------------------------------------------------ revestimentos
sub = {}
for p in D['pisos']:
    sub[p['cod']] = round(sub.get(p['cod'], 0) + p['area'], 3)
R = D['rod_total']
H_ROD = 0.08
rod1, rod2 = round(R['R01'] * H_ROD, 3), round(R['R02'] * H_ROD, 3)
W = DP['totais']
cx = lambda area, m2: math.ceil(area / m2 - 1e-9)

REV = [
    dict(cod='P01 · R01 · RP02', prod='Porcelanato Santorini OFW NAT', fab='Portinari', ref='A CONFIRMAR na loja', formato='90 × 90 cm · 7,0 mm · retificado',
         acab='Natural (NAT)', comp=[('Piso P01', sub['P01']), (f'Rodapé R01 {br(R["R01"])} m × 0,08', rod1), ('Parede RP02', W['RP02'])], m2cx=2.43,
         obs='Rodapé recortado das mesmas peças (perda de corte não incluída). Rejunte cor Corda, junta seca.'),
    dict(cod='P02 · R02', prod='Porcelanato Santorini SGR HARD (Stone Gray)', fab='Portinari', ref='62928', formato='90 × 90 cm · 8,0 mm · retificado',
         acab='HARD (antiderrapante)', comp=[('Piso P02 (varanda + deck)', sub['P02']), (f'Rodapé R02 {br(R["R02"])} m × 0,08', rod2)], m2cx=1.62,
         obs='Rejunte cor Corda, junta seca.'),
    dict(cod='P02 (opcional)', prod='Porcelanato Santorini SGR HARD — Corredor Lateral Externo', fab='Portinari', ref='62928', formato='90 × 90 cm · 8,0 mm',
         acab='HARD (antiderrapante)', comp=[('Corredor Lateral Externo', sub['NE'])], m2cx=1.62,
         obs='Somente se o corredor for em P02 (alternativa: piso intertravado drenante, fora da Bandini).'),
    dict(cod='RP03', prod='Porcelanato Confete WH NAT', fab='CEUSA', ref='5041209A', formato='100 × 100 cm · 9,0 mm · retificado', acab='Natural',
         comp=[('Banheiro Casal — fundo do box + nicho', W['RP03'])], m2cx=2.00, obs='Listado no site da Bandini (CONFET WH NAT 41209).'),
    dict(cod='RP04', prod='Porcelanato Confete PK NAT (Pink)', fab='CEUSA', ref='5041210A', formato='100 × 100 cm · 9,0 mm (A CONFIRMAR)', acab='Natural',
         comp=[('WC 01 e WC 02 — fundo do box + nicho', W['RP04'])], m2cx=2.00,
         obs='FORMATO A CONFIRMAR: a imagem do projeto mostra peças pequenas. m²/caixa conforme o 100 × 100.'),
    dict(cod='RP05', prod='Revestimento Fatto Oliva AC', fab='Decortiles', ref='SC 8068139', formato='30 × 90 cm', acab='Acetinado',
         comp=[('Banheiro Externo — fundo do box + nicho', W['RP05'])], m2cx=None, obs='m² por caixa A CONFIRMAR com a loja.'),
    dict(cod='RP07', prod='Revestimento Eliane Forma Branco AC (fundo da marcenaria)', fab='Eliane', ref='SC 8039383', formato='30 × 40 cm · 6,5 mm',
         acab='Acetinado', comp=[('Cozinha — paredes esquerda e superior', W['RP07'])], m2cx=1.92, obs='Junta 2 mm (fabricante).'),
]
for r in REV:
    r['area'] = round(sum(v for _, v in r['comp']), 3)
    r['cx'] = cx(r['area'], r['m2cx']) if r['m2cx'] else None

COMPL = [
    ('Rejunte cor Corda', 'P01, P02, R01, R02, RP02 (junta seca — fabricante Portinari)', 'Marca e tipo a definir (proposta: NBR 14992 tipo II em áreas externas e molhadas)'),
    ('Rejunte Confete Pink', 'RP04', 'Cores sugeridas pelo fabricante: Quartzolit Preto Grafite, Rejuntabras Marfim ou Quartzobras Castor'),
    ('Rejunte Confete WH, Fatto Oliva e Eliane Forma', 'RP03, RP05, RP07', 'Cor e tipo A CONFIRMAR'),
    ('Argamassa colante', 'Todos os revestimentos', 'Tipo e consumo conforme fabricante (porcelanato grande formato; áreas externas)'),
    ('Perfil metálico em L (cantoneira)', 'Desníveis entre pisos iguais (detalhe D2)', 'Sem desníveis documentados — quantidade A CONFIRMAR'),
]

FORA = [
    ('Pedra Moledo', 'Pedreira', f"{br(W['RP01'])} m²", 'Sala: fachadas sul 12,48 e leste 8,21 m² (h 3,20 m, desconta J01) + parede da TV 12,24 m² (h 3,40 m)'),
    ('Travertino Rock Face', 'Pedreira', f"{br(W['RP06'])} m²", 'Churrasqueira (frente e lateral) + faixa da bancada da varanda'),
    ('Pedra portuguesa branca', 'Fornecedor de pedras', f"{br(sub['PP01'])} m²", 'Rampa 37,00 + acesso externo 15,03 + calçada 98,07 m²'),
    ('Soleira mármore Itaúnas, baguete 5 cm', 'Marmoraria', f"{br(sum(s['comp'] for s in D['soleiras']))} m",
     '9 peças: ' + ' · '.join(f"{s['t']} {br(s['comp'])}" for s in D['soleiras'])),
    ('Piso intertravado drenante (opcional)', 'A definir', f"{br(sub['NE'])} m²", 'Somente se o Corredor Lateral Externo não for em P02'),
]

# ------------------------------------------------------------ louças e metais
AMB = [
    ('01', 'WC 01 e WC 02', 'Banho 1 (Suíte 1) e Banho 2 (Suíte 2) — mesma especificação; quantidades somam os dois'),
    ('02', 'Banheiro Casal', 'Banho 3 — Suíte Master'),
    ('03', 'Banheiro Externo', 'Banho 4 — WC Externo'),
    ('04', 'Cozinha', ''), ('05', 'Varanda Gourmet', 'Área Gourmet'), ('06', 'Lavanderia', 'Área de Serviço'), ('07', 'Área externa / piscina', ''),
]
KIT = ('Chuveiro', 'Kit chuveiro de parede com desviador e ducha manual', 'Deca / Acqua Plus', '—', 'Cromado')
ACB = ('Chuveiro', 'Acabamento de registro para misturador de duas alavancas (água quente e fria)', 'Deca / Level', '4900.C26.GD', 'Cromado')
DUC = ('Ducha higiênica', 'Ducha higiênica com registro e derivação', 'Deca / Level', '1984.C26.ACT', 'Cromado')
BAC = ('Bacia sanitária', 'Bacia sanitária com caixa acoplada, acionamento por botão duplo', 'Roca / ONA', '—', 'Padrão do fabricante')
ASS = ('Bacia sanitária', 'Assento sanitário original da linha', 'Roca / ONA', '—', 'Padrão do fabricante')
CUB = ('Lavatório', 'Cuba de embutir retangular 50×40 cm, louça', 'Deca', '—', 'Branca')
TOR = ('Lavatório', 'Torneira de mesa para lavatório, bica baixa', 'Deca / Level', '1193.C26', 'Cromado')
VAL = ('Lavatório', 'Válvula de escoamento tipo click para lavatório', 'Deca', '1601.C.CLI', 'Cromado')
ENG = ('Lavatório', 'Engate flexível com malha de aço inox trançado', 'Deca', '4607.C.030 ou 4607.C.040', 'Cromado')
BRS = ('Acessórios', 'Porta-toalha barra simples', 'Deca / You', '—', 'Cromado')
ROS = ('Acessórios', 'Porta-toalha de rosto / cabide', 'Deca / You', '—', 'Cromado')
PAP = ('Acessórios', 'Papeleira', 'Deca / You', '—', 'Cromado')
SIF = ('Sifão/engates', 'Kit sifão articulado 1½" + engates flexíveis para cuba de cozinha', 'Deca', '1682.C.112 (referência)', 'Cromado ou inox')
OBS_CX = 'Caixa acoplada: dispensa válvula de descarga; exige ponto de água fria na entrada da caixa.'
OBS_EN = 'Comprimento (30 ou 40 cm) a definir em obra.'
LM = [  # (ambiente, item, qtd, obs)
    ('01', KIT, 2, ''), ('01', ACB, 2, 'Compatível com o kit Acqua Plus.'), ('01', DUC, 2, ''), ('01', BAC, 2, OBS_CX), ('01', ASS, 2, ''),
    ('01', CUB, 2, 'Bancada com recorte para cuba de embutir.'), ('01', TOR, 2, 'Torneira convencional (definido).'), ('01', VAL, 2, ''),
    ('01', ENG, 2, OBS_EN), ('01', BRS, 2, ''), ('01', ROS, 2, ''), ('01', PAP, 2, ''),
    ('02', KIT, 1, ''), ('02', ACB, 1, 'Compatível com o kit Acqua Plus.'), ('02', DUC, 1, ''), ('02', BAC, 1, OBS_CX), ('02', ASS, 1, ''),
    ('02', CUB, '1 ou 2', 'Depende de a bancada ser simples ou dupla — A CONFIRMAR.'), ('02', TOR, '1 por cuba', ''), ('02', VAL, '1 por cuba', ''),
    ('02', ENG, '1 por cuba', OBS_EN), ('02', ('Acessórios', 'Porta-toalha barra dupla 60 cm', 'Deca / You', '2042.C104.060', 'Cromado'), 1, ''),
    ('02', ROS, 1, ''), ('02', PAP, 1, ''),
    ('03', KIT, 1, ''), ('03', ACB, 1, 'Compatível com o kit Acqua Plus.'), ('03', DUC, 1, ''), ('03', BAC, 1, OBS_CX), ('03', ASS, 1, ''),
    ('03', ('Lavatório', 'Cuba de apoio/sobrepor retangular, cerâmica', 'Deca / Slim', '—', 'Verde fosco'), 1,
     'Exige torneira de bica alta. Confirmar a cor verde fosco com a loja (foto em branco, representativa).'),
    ('03', ('Lavatório', 'Torneira de mesa bica alta para lavatório', 'Deca / Tube', '1198.GF.TUB.MT', 'Dark Antracite (fosco)'), 1,
     'Única peça neste acabamento. Confirmar se o código 1198 (bica alta) está disponível; foto da 1197 (bica baixa).'),
    ('03', VAL, 1, 'Não fabricada em Dark Antracite; mantido cromado (decisão do cliente).'), ('03', ENG, 1, OBS_EN),
    ('03', BRS, 1, ''), ('03', ROS, 1, ''), ('03', PAP, 1, ''),
    ('04', ('Cuba', 'Cuba de embutir retangular 75×40 cm, aço inox', 'Deca / Suprema', 'CC67075INX', 'Inox'), 1, ''),
    ('04', ('Torneira', 'Misturador monocomando de mesa para cozinha', 'Deca', '2275.INX', 'Inox'), 1, ''),
    ('04', ('Registro', 'Registro de gaveta até 1", acabamento quadrado', 'Deca', '4900.INX105.PQ', 'Inox'), 1, 'Instalado sob a bancada.'),
    ('04', SIF, 1, 'Referência exata compatível com CC67075INX (e versão inox) a confirmar com a loja.'),
    ('04', ('Filtro', 'Filtro/purificador de água com retenção de metais pesados', 'Purific / Camadas 10', '—', 'Torneira: escolher cor'), 1,
     'Acompanha torneira própria. Ponto de água a 1,00 m. Verificar se a Bandini revende.'),
    ('05', ('Cuba', 'Cuba de embutir retangular 50×40 cm, aço inox', 'Deca / Suprema', 'CC66050INX', 'Inox'), 1, ''),
    ('05', ('Torneira', 'Torneira de mesa bica móvel', 'Deca / Flex Plus', '1167.C21', 'Cromado'), 1, ''),
    ('05', ('Registro', 'Registro de gaveta, acabamento', 'Deca / Level', '4900.C26.GD', 'Cromado'), 1, 'Mesmo acabamento do restante da casa.'),
    ('05', SIF, 1, 'Referência exata compatível com CC66050INX a confirmar com a loja.'),
    ('06', ('Tanque', 'Tanque moldado 60×57 cm, 39 L, marmorizado sintético', 'I.Corso / Premium', '—', 'Branco'), 1, 'Verificar se a Bandini revende.'),
    ('06', ('Torneira', 'Torneira de parede bica longa para tanque', 'Deca / Link', '1178.C.LNK', 'Cromado'), 1, ''),
    ('06', ('Varal', 'Varal de parede dobrável tipo sanfona, 7 varetas, 120 cm', 'Varal Mágico', '—', 'Alumínio branco'), 1,
     'Fixação em parede com estrutura para o peso. Verificar se a Bandini revende.'),
    ('07', ('Ducha externa', 'Kit chuveirão de parede/piscina com acabamento de registro, só água fria', 'Deca', 'A CONFIRMAR', 'Cromado'), 1,
     'Referência exata do kit completo a fechar com a loja.'),
    ('07', ('Torneira de jardim', 'Torneira para jardim e tanque com adaptador para mangueira', 'Deca / Izy', '1153.C37', 'Cromado'), 1, ''),
]
assert len(LM) == 50
NAO_ORCAR = [
    'Box dos 4 banheiros: vidro temperado e ferragem (previsão: ferragem cromada).',
    'Cozinha: torneira independente do lava-louças (inox ou cromado, a confirmar).',
    'Lavanderia: torneira para 2º tanque ou ponto extra, se houver (sugestão: Deca Link 1178.C.LNK); registro/torneira da máquina de lavar (cromado).',
    'Varanda coberta: torneira para lavagem/manutenção, se houver ponto de água.',
    'Saboneteira: não haverá em nenhum banheiro. Lixeiras: fora do escopo (compra pelos clientes).',
]


def consolidado():
    agg = {}
    for amb, it, q, _ in LM:
        k = it[3] if it[3] not in ('—', 'A CONFIRMAR') else it[1] + ' · ' + it[2] + ' · ' + it[4]
        if k not in agg:
            agg[k] = (it, [])
        agg[k][1].append((amb, q))
    out = []
    for k, (it, lst) in agg.items():
        ref, desc, marca, cor = it[3], it[1], it[2], it[4]
        if ref == '4900.C26.GD':
            desc = 'Acabamento de registro (chuveiros dos banheiros e registro da Varanda Gourmet)'
        nums = [q for _, q in lst if isinstance(q, int)]
        var = [q for _, q in lst if not isinstance(q, int)]
        if var:  # itens do Banheiro Casal que dependem de 1 ou 2 cubas
            base = sum(nums)
            qtd = f'{base + 1} ou {base + 2}'
        else:
            qtd = sum(nums)
        onde = ', '.join(dict(AMB_N)[a] for a, _ in lst)
        out.append(dict(ref=ref, desc=desc, marca=marca, cor=cor, qtd=qtd, onde=onde))
    return out


AMB_N = [(a, n) for a, n, _ in AMB]

# ================================================================== HTML
CSS = """
@page { size: 297mm 210mm; margin: 12mm 12mm 16mm 12mm; }
* { box-sizing: border-box; margin: 0; padding: 0; }
body { font-family: 'Manrope'; color: #22251a; font-size: 7.4pt; background: #fff; -webkit-print-color-adjust: exact; print-color-adjust: exact; }
.foot { position: fixed; bottom: -10mm; left: 0; right: 0; display: flex; justify-content: space-between; font-size: 6.3pt; color: #4f5439; border-top: 0.25mm solid #d8cfbc; padding-top: 2mm; }
.hero { display: flex; gap: 5mm; margin-bottom: 4mm; }
.hero .a { flex: 1; background: #903f2a; color: #f5e7c4; border-radius: 0 9mm 0 0; padding: 7mm 8mm; }
.hero .a .k { font-size: 6.6pt; letter-spacing: 2.4pt; font-weight: 600; }
.hero .a h1 { font-family: 'Outfit'; font-weight: 400; font-size: 25pt; line-height: 1.1; margin: 2.2mm 0; }
.hero .a p { font-size: 8.4pt; opacity: .95; }
.hero .b { width: 78mm; background: #ece4d4; border-radius: 0 0 0 9mm; padding: 6mm 7mm; text-align: right; font-size: 7.4pt; }
.hero .b img { height: 8mm; margin-bottom: 7mm; filter: brightness(.32); }
h2 { display: flex; align-items: center; gap: 3.5mm; font-family: 'Outfit'; font-weight: 400; font-size: 15pt; margin: 6mm 0 2.5mm; break-after: avoid; }
h2 .n { background: #3f422f; color: #f5e7c4; font-family: 'Manrope'; font-size: 8pt; font-weight: 700; border-radius: 2mm; padding: 1.6mm 2.4mm; }
h2 small { font-family: 'Manrope'; font-size: 7pt; color: #4f5439; }
table { width: 100%; border-collapse: separate; border-spacing: 0; }
thead { display: table-header-group; }
th { background: #3f422f; color: #f5e7c4; font-size: 6.2pt; font-weight: 700; text-transform: uppercase; letter-spacing: .3pt; text-align: left; padding: 2mm 1.6mm; }
th:first-child { border-top-left-radius: 2.2mm; } th:last-child { border-top-right-radius: 2.2mm; }
td { padding: 1.3mm 1.6mm; border-bottom: 0.25mm solid #d8cfbc; vertical-align: top; line-height: 1.32; }
tr { break-inside: avoid; }
td.n, th.n { text-align: right; white-space: nowrap; }
td.ref { color: #903f2a; font-weight: 700; }
td.p { background: #fbf8f1; border-left: 0.25mm dashed #cbbfa5; min-width: 17mm; }
tr.sec td { background: #ece4d4; color: #7a3423; font-weight: 700; letter-spacing: 1.6pt; font-size: 6.4pt; }
tr.tot td { font-weight: 700; font-size: 7.8pt; border-bottom: none; }
.foto { width: 13mm; height: 13mm; object-fit: contain; background: #fff; border-radius: 1.6mm; border: 0.25mm solid #e1d9c8; display: block; }
.note { font-size: 6.6pt; color: #4f5439; margin-top: 1.6mm; line-height: 1.4; }
.box { background: #f5e7c4; border-radius: 3mm; padding: 4mm 5mm; margin-top: 3mm; }
.box li { margin: 0 0 1.2mm 4mm; line-height: 1.4; }
.grid { display: grid; grid-template-columns: 1.2fr 1fr; gap: 5mm; }
.pg { break-before: page; }
.warn { color: #7a3423; }
"""


def fonts():
    s = ''
    for w in (400, 500, 600, 700):
        s += f"@font-face{{font-family:'Manrope';font-weight:{w};src:url('fonts/manrope-latin-{w}-normal.woff2')}}"
    for w in (300, 400):
        s += f"@font-face{{font-family:'Outfit';font-weight:{w};src:url('fonts/outfit-latin-{w}-normal.woff2')}}"
    return s


def preco_cols():
    return "<td class='p'></td><td class='p'></td>"


def html_doc():
    h = ''
    h += f"""<div class="hero"><div class="a"><div class="k">LISTA TÉCNICA PARA ORÇAMENTO · LOJA BANDINI</div>
<h1>Residência IC — Ivan e Carol</h1><p>Revestimentos de piso e parede, louças e metais — quantitativos para orçamento e compra na mesma loja.</p></div>
<div class="b"><img src="logo_duas.png"><div><b>Data:</b> {DATA} · <b>Revisão</b> {REVN}</div><div style="margin-top:1mm">Lidiane Barbosa e Isadora Ferrari</div></div></div>
<div class="grid"><div>
<div style="font-size:8pt;line-height:1.5">Documento para a loja preencher preço unitário, total e prazo em cada linha (colunas tracejadas). Revestimentos pelos
Cadernos de Detalhamento 01/02 (piso) e 02/02 (parede), Rev. 00; louças e metais pela Lista técnica de compras de 25/09/2026. O modelo vale pela referência/código.</div>
<div class="box"><b>Condições para o orçamento</b><ul style="margin-top:1.6mm">
<li>Quantidades <b>líquidas</b>, sem perdas de compra ou de corte; caixas arredondadas para cima. A perda é definida pela arquitetura/loja.</li>
<li>Informar disponibilidade, lote/tonalidade (comprar cada revestimento do mesmo lote) e prazo de entrega.</li>
<li>Itens sem código ou marcados “A CONFIRMAR”: a loja indica a referência exata antes de fechar o pedido.</li>
<li>Itens de pedreira, marmoraria e opções fora da Bandini estão na seção 04, só para controle.</li></ul></div>
</div><div>{resumo()}</div></div>"""
    # 01 revestimentos
    rows = ''
    for r in REV:
        comp = '<br>'.join(f"{esc(n)}: {br(v)}" for n, v in r['comp'])
        rows += (f"<tr><td class='ref'>{esc(r['cod'])}</td><td><b>{esc(r['prod'])}</b><div class='note'>{esc(r['formato'])} · {esc(r['acab'])}</div></td>"
                 f"<td>{esc(r['fab'])}<div class='note'>ref. {esc(r['ref'])}</div></td><td style='font-size:6.6pt'>{comp}</td>"
                 f"<td class='n'><b>{br(r['area'])}</b></td><td class='n'>{br(r['m2cx']) if r['m2cx'] else 'A CONFIRMAR'}</td>"
                 f"<td class='n'><b>{r['cx'] if r['cx'] else '—'}</b></td>{preco_cols()}<td style='font-size:6.5pt'>{esc(r['obs'])}</td></tr>")
    h += f"""<h2><span class="n">01</span>Revestimentos de piso e parede <small>quantidades líquidas · preço por caixa</small></h2>
<table><thead><tr><th>Cód.</th><th>Produto</th><th>Fabricante</th><th>Composição (m²)</th><th class="n">m² líq.</th><th class="n">m²/cx</th><th class="n">Caixas</th>
<th>Preço/cx (R$)</th><th>Total (R$)</th><th>Observação</th></tr></thead><tbody>{rows}
<tr class='tot'><td colspan='7'>Total revestimentos (sem a linha opcional)</td><td class='p'></td><td class='p'></td><td></td></tr></tbody></table>
<div class="note">Caixas = m² líquidos ÷ m² por caixa, arredondado para cima. Se o Corredor Lateral Externo for em P02, somar a linha opcional ao P02 (total {br(REV[1]['area'] + REV[2]['area'])} m² = {cx(REV[1]['area'] + REV[2]['area'], 1.62)} caixas).
Se o rodapé sair dos trechos revestidos até o piso (pendência do Caderno 02/02), o P01 reduz até {br(28.35 * 0.08)} m².</div>"""
    rows = ''.join(f"<tr><td><b>{esc(a)}</b></td><td>{esc(b)}</td><td style='font-size:6.6pt'>{esc(c)}</td><td class='n'>A CONFIRMAR</td>{preco_cols()}</tr>" for a, b, c in COMPL)
    h += f"""<h2><span class="n">02</span>Complementos de assentamento <small>quantidade conforme consumo do fabricante</small></h2>
<table><thead><tr><th>Item</th><th>Onde</th><th>Especificação</th><th class="n">Qtd.</th><th>Preço unit. (R$)</th><th>Total (R$)</th></tr></thead><tbody>{rows}</tbody></table>"""
    # 03 louças e metais por ambiente
    h += '<div class="pg"></div>'
    h += """<h2><span class="n">03</span>Louças e metais por ambiente <small>conforme a Lista técnica de compras de 25/09/2026</small></h2>"""
    for i, (a, nome, sub_) in enumerate(AMB):
        rows = ''
        for k, (amb, it, q, obs) in enumerate(LM):
            if amb != a:
                continue
            rows += (f"<tr><td><img class='foto' src='imagens/loucas/f{k + 1:02d}.png'></td><td><b>{esc(it[0])}</b></td><td>{esc(it[1])}</td><td>{esc(it[2])}</td>"
                     f"<td class='ref'>{esc(it[3])}</td><td>{esc(it[4])}</td><td class='n'><b>{esc(q)}</b></td>{preco_cols()}<td style='font-size:6.5pt'>{esc(obs)}</td></tr>")
        h += f"""<table style="table-layout:fixed;margin-top:{'0' if i == 0 else '4mm'}{';break-inside:avoid' if rows.count('<tr>') <= 6 else ''}"><thead><tr><th colspan="10" style="background:#ece4d4;color:#7a3423;letter-spacing:1.6pt">{a} · {esc(nome).upper()}{(' — ' + esc(sub_)) if sub_ else ''}</th></tr>
<colgroup><col style="width:17mm"><col style="width:22mm"><col style="width:62mm"><col style="width:26mm"><col style="width:30mm"><col style="width:24mm"><col style="width:15mm"><col style="width:20mm"><col style="width:20mm"><col></colgroup>
<tr><th style="border-radius:0">Foto</th><th>Item</th><th>Descrição técnica</th><th>Marca/linha</th><th>Referência</th><th>Cor/acab.</th><th class="n">Qtd.</th><th>Preço unit. (R$)</th><th>Total (R$)</th><th style="border-radius:0">Observação</th></tr></thead>
<tbody>{rows}</tbody></table>"""
    # consolidado
    rows = ''.join(f"<tr><td class='ref'>{esc(c['ref'])}</td><td>{esc(c['desc'])}</td><td>{esc(c['marca'])}</td><td>{esc(c['cor'])}</td>"
                   f"<td class='n'><b>{esc(c['qtd'])}</b></td><td style='font-size:6.5pt'>{esc(c['onde'])}</td>{preco_cols()}</tr>" for c in consolidado())
    h += f"""<div class="pg"></div><h2><span class="n">03A</span>Louças e metais — consolidado por referência <small>para o pedido na loja</small></h2>
<table><thead><tr><th>Referência</th><th>Descrição</th><th>Marca/linha</th><th>Cor/acab.</th><th class="n">Qtd. total</th><th>Ambientes</th><th>Preço unit. (R$)</th><th>Total (R$)</th></tr></thead>
<tbody>{rows}<tr class='tot'><td colspan='6'>Total louças e metais</td><td class='p'></td><td class='p'></td></tr></tbody></table>
<div class="note">Onde há “1 ou 2”: depende de a bancada do Banheiro Casal ter uma ou duas cubas (A CONFIRMAR).</div>"""
    rows = ''.join(f"<tr><td><b>{esc(a)}</b></td><td>{esc(b)}</td><td class='n'><b>{esc(c)}</b></td><td style='font-size:6.6pt'>{esc(d)}</td><td class='p'></td></tr>" for a, b, c, d in FORA)
    h += f"""<h2><span class="n">04</span>Fora da Bandini — controle <small>pedreira, marmoraria e opções</small></h2>
<table><thead><tr><th>Item</th><th>Fornecedor</th><th class="n">Qtd. líquida</th><th>Onde</th><th>Orçamento (R$)</th></tr></thead><tbody>{rows}</tbody></table>
<div class="grid" style="margin-top:3mm"><div class="box"><b>Não orçar por enquanto (em definição)</b><ul style="margin-top:1.6mm">{''.join(f'<li>{esc(x)}</li>' for x in NAO_ORCAR)}</ul></div>
<div class="box"><b>Confirmar com a loja</b><ul style="margin-top:1.6mm">
<li>Santorini OFW NAT e SGR HARD (Portinari): disponibilidade na Bandini (não localizados no site; a loja trabalha com Portinari).</li>
<li>Confete Pink: formato (100 × 100 cm ou peça menor, conforme imagem do projeto).</li>
<li>Fatto Oliva: m² por caixa. Cuba Slim verde fosco e torneira Tube 1198 (bica alta).</li>
<li>Kits de sifão compatíveis com as cubas Suprema e kit completo do chuveirão externo.</li>
<li>Purific, I.Corso e Varal Mágico: se a Bandini revende.</li></ul></div></div>"""
    foot = f"<div class='foot'><span><b>DUAS</b> design & arq · Lidiane Barbosa e Isadora Ferrari</span><span>Residência IC — Lista técnica para orçamento (Bandini) · {DATA} · Rev. {REVN}</span></div>"
    return f"<!doctype html><html lang='pt-BR'><head><meta charset='utf-8'><title>Lista técnica para orçamento — Residência IC — Rev. {REVN}</title><style>{fonts()}{CSS}</style></head><body>{h}</body></html>"


def resumo():
    n_rev = len(REV)
    n_lm = len(LM)
    return f"""<table><thead><tr><th>Grupo</th><th class="n">Linhas</th><th>Fornecedor</th></tr></thead><tbody>
<tr><td>01 Revestimentos (Portinari, CEUSA, Decortiles, Eliane)</td><td class="n">{n_rev}</td><td>Bandini</td></tr>
<tr><td>02 Complementos (rejuntes, argamassa, perfil)</td><td class="n">{len(COMPL)}</td><td>Bandini</td></tr>
<tr><td>03 Louças e metais (Deca, Roca e outros) — 7 ambientes</td><td class="n">{n_lm}</td><td>Bandini</td></tr>
<tr><td>04 Pedras, soleiras e opções</td><td class="n">{len(FORA)}</td><td>Pedreira / marmoraria</td></tr>
</tbody></table>
<div class="note">Revestimentos: P01 {REV[0]['cx']} cx · P02 {REV[1]['cx']} cx · RP03 {REV[3]['cx']} cx · RP04 {REV[4]['cx']} cx · RP07 {REV[6]['cx']} cx · RP05 A CONFIRMAR.</div>"""


# ================================================================== XLSX
def planilha():
    wb = Workbook()
    hdr = PatternFill('solid', fgColor='3F422F')
    sec = PatternFill('solid', fgColor='ECE4D4')
    inp = PatternFill('solid', fgColor='FBF3DF')

    def head(ws, cols, widths):
        ws.append(cols)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + i)].width = w
            c = ws.cell(row=ws.max_row, column=i); c.fill = hdr; c.font = Font(bold=True, color='F5E7C4'); c.alignment = Alignment(wrap_text=True, vertical='center')
        ws.freeze_panes = 'A2'

    ws = wb.active; ws.title = 'Revestimentos'
    head(ws, ['Código', 'Produto', 'Fabricante', 'Ref.', 'Formato', 'Composição', 'm² líquidos', 'm²/caixa', 'Caixas', 'Preço por caixa (R$)', 'Total (R$)', 'Observação'],
         [16, 40, 12, 14, 26, 46, 11, 10, 9, 14, 14, 50])
    for r in REV:
        i = ws.max_row + 1
        ws.append([r['cod'], r['prod'], r['fab'], r['ref'], r['formato'], '; '.join(f'{n}: {v:.2f}' for n, v in r['comp']), r['area'], r['m2cx'] or 'A CONFIRMAR',
                   f'=IF(ISNUMBER(H{i}),ROUNDUP(G{i}/H{i},0),"")', None, f'=IF(AND(ISNUMBER(I{i}),ISNUMBER(J{i})),I{i}*J{i},"")', r['obs']])
        ws.cell(row=i, column=10).fill = inp
    last = ws.max_row
    ws.append(['', 'TOTAL (sem a linha opcional)'] + [''] * 8 + [f'=SUM(K2:K3)+SUM(K5:K{last})'])
    ws.cell(row=ws.max_row, column=2).font = Font(bold=True)
    ws.append([])
    ws.append(['Complementos'] + [''] * 11)
    for c in ws[ws.max_row]: c.fill = sec
    for a, b, c in COMPL:
        ws.append([a, b, '', '', c, '', 'A CONFIRMAR'])

    ws = wb.create_sheet('Louças e metais')
    head(ws, ['Ambiente', 'Item', 'Descrição técnica', 'Marca/linha', 'Referência', 'Cor/acabamento', 'Qtd.', 'Preço unit. (R$)', 'Total (R$)', 'Observação'],
         [22, 16, 48, 18, 22, 18, 10, 14, 14, 60])
    nomes = dict(AMB_N)
    for amb, it, q, obs in LM:
        i = ws.max_row + 1
        ws.append([nomes[amb], it[0], it[1], it[2], it[3], it[4], q, None, f'=IF(AND(ISNUMBER(G{i}),ISNUMBER(H{i})),G{i}*H{i},"")', obs])
        ws.cell(row=i, column=8).fill = inp
    ws.append(['', 'TOTAL'] + [''] * 6 + [f'=SUM(I2:I{ws.max_row})'])

    ws = wb.create_sheet('Consolidado Bandini')
    head(ws, ['Referência', 'Descrição', 'Marca/linha', 'Cor/acabamento', 'Qtd. total', 'Ambientes', 'Preço unit. (R$)', 'Total (R$)'], [24, 50, 18, 18, 11, 40, 14, 14])
    for c in consolidado():
        i = ws.max_row + 1
        ws.append([c['ref'], c['desc'], c['marca'], c['cor'], c['qtd'], c['onde'], None, f'=IF(AND(ISNUMBER(E{i}),ISNUMBER(G{i})),E{i}*G{i},"")'])
        ws.cell(row=i, column=7).fill = inp
    ws.append(['', 'TOTAL'] + [''] * 5 + [f'=SUM(H2:H{ws.max_row})'])

    ws = wb.create_sheet('Fora da Bandini')
    head(ws, ['Item', 'Fornecedor', 'Qtd. líquida', 'Onde', 'Orçamento (R$)'], [36, 22, 14, 70, 16])
    for a, b, c, d in FORA:
        ws.append([a, b, c, d, None])
    ws = wb.create_sheet('Não orçar')
    for x in NAO_ORCAR:
        ws.append([x])
    ws.column_dimensions['A'].width = 120
    for s in wb.worksheets:
        for row in s.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical='top')
    wb.save(os.path.join(AQUI, '..', 'quantitativos', 'ORCAMENTO BANDINI - REVESTIMENTOS LOUCAS E METAIS - R00.xlsx'))


if __name__ == '__main__':
    open(os.path.join(AQUI, 'orcamento.html'), 'w').write(html_doc())
    planilha()
    for r in REV:
        print(r['cod'], r['area'], r['m2cx'], r['cx'])
    for c in consolidado():
        print(c['ref'], '|', c['qtd'], '|', c['onde'])
